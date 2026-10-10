import logging

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy.orm import Session, joinedload

# Giả định bạn đã có hàm get_db cung cấp SQLAlchemy Session
from app.db.base import get_db

# Import các ORM model
from app.db.models import Attachment, Email, EmailLabel, Label

router = APIRouter()
logger = logging.getLogger(__name__)

# ==========================================
# 1. PYDANTIC SCHEMAS (Request/Response)
# ==========================================
class LabelAssignRequest(BaseModel):
    label_id: str  # nhận label_name (VD: "PO", "QUOTE") từ frontend
    reason: str | None = "Người dùng gán thủ công từ UI"

# ==========================================
# 2. ENDPOINTS
# ==========================================

@router.get("/")
def get_emails(
    status: str | None = None,
    limit: int = Query(20, ge=1, le=100),
    offset: int = Query(0, ge=0),
    db: Session = Depends(get_db)
):
    """Lấy danh sách email cho Smart Inbox (Hiển thị thẻ tóm tắt siêu tốc)"""
    # Dùng joinedload để tránh lỗi N+1 query khi kéo thông tin Contact và Summary
    query = db.query(Email).options(
        joinedload(Email.sender),
        joinedload(Email.summary)
    )

    if status:
        query = query.filter(Email.status == status)

    query = query.order_by(Email.received_at.desc())
    total = query.count()
    emails = query.offset(offset).limit(limit).all()

    # Flatten dữ liệu để Frontend dễ render List/Table
    results = []
    for e in emails:
        results.append({
            "id_email": str(e.id_email),
            "subject": e.subject,
            "sender_name": e.sender.sender_name if e.sender else "Unknown",
            "sender_email": e.sender.email_address if e.sender else "Unknown",
            "status": e.status,
            "received_at": e.received_at.isoformat() if e.received_at else None,
            "tldr": e.summary.content_summarized if e.summary else "",
            "security_flags": e.security_flags # Trả về luôn chuỗi JSON cảnh báo để UI hiện màu đỏ
        })

    return {"total": total, "items": results}


@router.get("/{email_id}")
def get_email_detail(email_id: str, db: Session = Depends(get_db)):
    """Lấy toàn bộ chi tiết của 1 email (Nội dung gốc, File đính kèm, Action Items)"""
    email = db.query(Email).options(
        joinedload(Email.sender),
        joinedload(Email.summary),
        joinedload(Email.attachments),
        # Nếu đã có nhãn thì join để lấy tên nhãn
        joinedload(Email.labels).joinedload(EmailLabel.label)
    ).filter(Email.id_email == email_id).first()

    if not email:
        raise HTTPException(status_code=404, detail="Không tìm thấy Email")

    return {
        "id_email": str(email.id_email),
        "id_mailbox": str(email.id_mailbox),
        "metadata": {
            "subject": email.subject,
            "received_at": email.received_at.isoformat(),
            "language": email.language,
            "status": email.status
        },
        "contact": {
            "id_sender": str(email.id_sender) if email.id_sender else None,
            "name": email.sender.sender_name if email.sender else None,
            "email": email.sender.email_address if email.sender else None,
            "cc_emails": email.cc_emails or [],
            "bcc_emails": email.bcc_emails or [],
            "company": email.sender.company if email.sender else None,
            "phone": email.sender.phone if email.sender else None,
            "is_internal": email.sender.is_internal if email.sender else False,
            "is_known_customer": email.sender.is_known_customer if email.sender else False
        },
        "content": {
            "message_id": email.message_id,
            "thread_id": email.thread_id,
            "in_reply_to": email.in_reply_to,
            "body_clean": email.body_clean,
            "body_full": email.body_text,
            "content_summarized": email.summary.content_summarized if email.summary else "",
            "received_at": email.received_at.isoformat(),
            "has_attachments": email.has_attachments
        },
        "ai_analysis": {
            "tldr": email.summary.content_summarized if email.summary else "",
            "key_points": email.summary.key_points if email.summary else {},
            "security": email.security_flags
        },
        "attachments": [
            {
                "id": str(att.id_attachment),
                "file_name": att.file_name,
                "size_bytes": att.size_bytes,
                "mime_type": att.mime_type,
                "process_status": att.process_status,
                "storage_key": att.storage_key,
                "is_inline": att.is_inline
            } for att in email.attachments
        ],
        "assigned_labels": [
            {
                "label_name": lbl.label.label_name,
                "source": lbl.source
            } for lbl in email.labels
        ]
    }


@router.get("/{email_id}/thread")
def get_email_thread(email_id: str, db: Session = Depends(get_db)):
    """Toàn bộ lịch sử hội thoại (thread) chứa email này — để UI hiển thị luồng mail dài"""
    email = db.query(Email).filter(Email.id_email == email_id).first()
    if not email:
        raise HTTPException(status_code=404, detail="Không tìm thấy Email")

    if not email.thread_id:
        return {"thread_id": None, "items": [_email_brief(e) for e in [email]]}

    rows = db.query(Email).options(
        joinedload(Email.sender),
        joinedload(Email.summary),
    ).filter(Email.thread_id == email.thread_id).order_by(Email.sent_at.asc()).all()

    return {"thread_id": email.thread_id, "items": [_email_brief(e) for e in rows]}


def _email_brief(e: Email) -> dict:
    return {
        "id_email": str(e.id_email),
        "subject": e.subject,
        "sender_name": e.sender.sender_name if e.sender else None,
        "sender_email": e.sender.email_address if e.sender else None,
        "sent_at": e.sent_at.isoformat() if e.sent_at else None,
        "received_at": e.received_at.isoformat() if e.received_at else None,
        "body_clean": e.body_clean,
        "content_summarized": e.summary.content_summarized if e.summary else "",
        "has_attachments": e.has_attachments,
    }


@router.post("/{email_id}/labels")
def assign_email_label(email_id: str, payload: LabelAssignRequest, db: Session = Depends(get_db)):
    """API Human-in-the-loop: Gán nhãn và kích hoạt luồng xử lý"""
    email = db.query(Email).filter(Email.id_email == email_id).first()
    if not email:
        raise HTTPException(status_code=404, detail="Không tìm thấy Email")

    label = db.query(Label).filter(Label.label_name == payload.label_id).first()
    if not label:
        raise HTTPException(status_code=404, detail="Nhãn không hợp lệ")

    # 1. Ghi nhận lịch sử gán nhãn thủ công — idempotency + xử lý is_primary
    existing = db.query(EmailLabel).filter(
        EmailLabel.id_email == email.id_email,
        EmailLabel.id_label == label.id_label,
        EmailLabel.source == "user",
    ).first()
    if existing:
        existing.is_primary = True
        existing.confidence = 1.0
        if payload.reason is not None:
            existing.reason = payload.reason
    else:
        # bỏ is_primary của mọi nhãn khác trên cùng email (kể cả nhãn AI) — HITL thắng
        db.query(EmailLabel).filter(
            EmailLabel.id_email == email.id_email,
            EmailLabel.is_primary.is_(True),
        ).update({EmailLabel.is_primary: False}, synchronize_session=False)
        new_email_label = EmailLabel(
            id_email=email.id_email,
            id_label=label.id_label,
            source="user",
            is_primary=True,
            confidence=1.0,
            reason=payload.reason,
        )
        db.add(new_email_label)

    # 2. Xử lý logic định tuyến (Routing) dựa trên loại nhãn
    if label.label_name == "SPAM_ADS":
        # Luồng 1: Nếu là quảng cáo -> Bỏ qua, kết thúc luồng.
        email.status = "IGNORED"
        email.assigned_agent = "NONE"
        message = "Đã đánh dấu là Quảng cáo. Bỏ qua email này."

    elif label.label_name == "EXCEPTION" or label.label_name == "SED":
        # Luồng 2: Trường hợp phức tạp / Tờ khai Mỹ -> Chuyển thẳng cho Human Manager (Không dùng AI)
        email.status = "ESCALATED"
        email.assigned_agent = "HUMAN_MANAGER"
        message = f"Trường hợp phức tạp ({label.label_name}). Đã chuyển cho bộ phận quản lý."

    else:
        # Luồng 3: Các nghiệp vụ chuẩn (PO, Quote, SOA, Proforma Invoice...)
        # Chuyển trạng thái sang PROCESSING để background worker đánh thức Agent tương ứng
        email.status = "PROCESSING"
        email.assigned_agent = f"{label.label_name}_AGENT"
        message = f"Đã gán nhãn {label.label_name}. Đang kích hoạt {email.assigned_agent} soạn thảo phản hồi..."

    db.commit()

    return {
        "status": "success",
        "assigned_agent": email.assigned_agent,
        "email_status": email.status,
        "message": message
    }


def _extract_doc_preview(data: bytes, mime: str, file_name: str, name_lower: str) -> tuple[list[list], str | None]:
    """Trích bảng (rows) cho preview: xlsx sheet đầu, docx tables. Trả (rows, error)."""
    if name_lower.endswith(".xlsx") or "spreadsheet" in mime or "excel" in mime:
        import io as _io

        from openpyxl import load_workbook
        wb = load_workbook(filename=_io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows: list[list] = []
        for row in ws.iter_rows(max_row=21, max_col=11, values_only=True):
            rows.append(["" if v is None else str(v) for v in row])
        wb.close()
        return [r for r in rows if any(c.strip() for c in r)], None
    if name_lower.endswith(".docx") or "wordprocessingml" in mime or mime == "application/msword":
        import io as _io

        from docx import Document
        doc = Document(_io.BytesIO(data))
        out: list[list] = []
        for para in doc.paragraphs[:15]:
            t = para.text.strip()
            if t:
                out.append([t])
        for table in doc.tables[:3]:
            for row in table.rows:
                out.append([cell.text.strip() for cell in row.cells])
        return out[:60], None
    return [], "Không hỗ trợ preview cho loại file này"


@router.get("/attachments/{attachment_id}")
def get_attachment_file(attachment_id: str, db: Session = Depends(get_db)):
    """Tải/preview file đính kèm — dùng cho iframe/pdf viewer trên UI"""
    import mimetypes as mt

    from app.services.storage import get_storage

    att = db.query(Attachment).filter(Attachment.id_attachment == attachment_id).first()
    if not att:
        raise HTTPException(status_code=404, detail="Không tìm thấy file")

    mime, _ = mt.guess_type(att.file_name)
    mime = att.mime_type or mime or "application/octet-stream"

    try:
        data = get_storage().get_file(att.storage_key)
        # Chuyển bytes sang response; inline disposition cho preview
        return Response(
            content=data,
            media_type=mime,
            headers={"Content-Disposition": f'inline; filename="{att.file_name}"'},
        )
    except Exception as e:
        logger.exception("get_file failed: %s", e)
        raise HTTPException(status_code=500, detail="Không tải được file") from e


@router.get("/attachments/{attachment_id}/preview")
def get_attachment_preview(attachment_id: str, db: Session = Depends(get_db)):
    """Trích nội dung text/bảng của docx/xlsx để hiển thị trực tiếp trên UI."""
    import mimetypes as mt

    from app.services.storage import get_storage

    att = db.query(Attachment).filter(Attachment.id_attachment == attachment_id).first()
    if not att:
        raise HTTPException(status_code=404, detail="Không tìm thấy file")

    mime, _ = mt.guess_type(att.file_name)
    mime = (att.mime_type or mime or "").lower()
    name_lower = att.file_name.lower()

    try:
        data = get_storage().get_file(att.storage_key)
    except Exception as e:
        logger.exception("get_file failed: %s", e)
        raise HTTPException(status_code=500, detail="Không tải được file") from e

    try:
        rows, err = _extract_doc_preview(data, mime, att.file_name, name_lower)
    except Exception as e:
        logger.exception("preview extraction failed: %s", e)
        rows, err = [], f"Không đọc được file: {e}"

    return {"file_name": att.file_name, "kind": "table", "rows": rows, "error": err}
